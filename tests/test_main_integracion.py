import openpyxl

import main
from tests.conftest import construir_linea


def _linea_0911():
    return construir_linea({
        "TIPO-REGISTRO": "01", "RED-LOGICA": "0911", "NUMERO-TARJETA": "1111",
        "FIID SPONSOR": "0013", "TIPO DE MENSAJE": "0210", "COD-TIPO-TRANS": "10",
        " CODIGO-RESP": "000", "MONTO-1": "0000001000000",
        "NUMERO -APROBACION": "AAA1", "INDICADOR INTER/NACIONAL": "I",
    })


def test_flujo_completo_genera_excel_con_3_hojas(tmp_path, monkeypatch):
    ruta = tmp_path / "GOF.GRB.FM14.F260813.txt"
    ruta.write_text(_linea_0911() + "\n", encoding="utf-8")

    # main.py hace `from config import REPORTE_EP`, asi que el nombre vive
    # como atributo propio de `main` (no de `config`) — hay que parchear
    # `main.REPORTE_EP` para que `opcion_generar_reporte()` (que referencia
    # el nombre importado, no `config.REPORTE_EP`) realmente lo use.
    # Esto evita que el test escriba sobre el archivo real de salidas/
    # (el reporte diario real del operador), que podria estar abierto en
    # Excel y hacer fallar el test con PermissionError por razones ajenas
    # al codigo bajo prueba.
    reporte_tmp = tmp_path / "EP-test.xlsx"
    monkeypatch.setattr("main.REPORTE_EP", reporte_tmp)

    main.estado["rutas"] = [ruta]

    monkeypatch.setattr("main.pausar", lambda: None)
    main.opcion_generar_reporte()

    assert reporte_tmp.exists()
    libro = openpyxl.load_workbook(reporte_tmp)
    assert len(libro.sheetnames) == 3
    assert "EFECTO CERO" in libro.sheetnames
