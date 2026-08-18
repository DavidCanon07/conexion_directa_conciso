"""
exportador.py
Escribe DataFrames a .xlsx con formato legible: encabezado resaltado,
ancho de columna automatico, encabezado congelado y formato de moneda
segun el nombre del campo (mismo patron que "conexion directa").

Las 3 hojas de este reporte son siempre DataFrames (nunca una hoja de
posicion libre como la "Hoja1" de Archivo 3 en "conexion directa"), asi
que se escriben directo con Workbook(write_only=True): openpyxl
transmite cada fila al XML de salida sin retener los objetos Cell en
memoria, evitando la degradacion no lineal que sufre un Workbook normal
con hojas de cientos de miles de filas (y ahora que LAYOUT trae los 55
campos completos, las hojas de salida son mas anchas que antes).
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils import logger, manejar_errores

_ANCHO_MAXIMO_COLUMNA = 40
_FILAS_MUESTRA_ANCHO = 2000  # suficiente para estimar el ancho de columnas de texto sin recorrer todo el DataFrame


def _calcular_formato_y_ancho(df: pd.DataFrame) -> tuple:
    """Formato numerico (moneda por substring del nombre) y ancho de columna, por columna."""
    formato_por_columna = {}
    ancho_por_columna = {}
    muestra = df.head(_FILAS_MUESTRA_ANCHO)

    for col_idx, nombre in enumerate(df.columns, start=1):
        if "monto" in str(nombre).lower() or "valor" in str(nombre).lower():
            formato_por_columna[col_idx] = "#,##0.00"
            ancho_por_columna[col_idx] = 18
        else:
            largo_muestra = muestra[nombre].astype(str).map(len).max() if len(muestra) else 0
            ancho_por_columna[col_idx] = min(max(largo_muestra, len(str(nombre))) + 3, _ANCHO_MAXIMO_COLUMNA)

    return formato_por_columna, ancho_por_columna


def _formatear_hoja(ws, df: pd.DataFrame):
    """
    Escribe encabezado (resaltado + centrado), congela la fila 1 y
    escribe los datos fila por fila con ws.append() en vez de
    DataFrame.to_excel() (que asigna celda por celda). En modo
    write_only cada fila se serializa al escribirla y ya no se puede
    volver a tocar, asi que el encabezado y el formato de moneda se
    arman ANTES de cada append con WriteOnlyCell; freeze_panes y
    column_dimensions tambien deben fijarse antes del primer append
    (verificado en openpyxl 3.1.5: despues, se ignoran silenciosamente).
    """
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center")

    columnas = list(df.columns)
    formato_por_columna, ancho_por_columna = _calcular_formato_y_ancho(df)

    ws.freeze_panes = "A2"
    for col_idx, ancho in ancho_por_columna.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho

    encabezado = []
    for nombre in columnas:
        celda = WriteOnlyCell(ws, value=nombre)
        celda.fill = header_fill
        celda.font = header_font
        celda.alignment = header_alignment
        encabezado.append(celda)
    ws.append(encabezado)

    # Excel/openpyxl no aceptan NaN/NaT; to_excel los convertia a celda
    # vacia automaticamente, aqui se reemplazan explicitamente por None.
    datos = df.where(pd.notnull(df), None)

    if not formato_por_columna:
        for fila in datos.itertuples(index=False, name=None):
            ws.append(fila)
        return

    for fila in datos.itertuples(index=False, name=None):
        fila_salida = list(fila)
        for col_idx, formato in formato_por_columna.items():
            celda = WriteOnlyCell(ws, value=fila_salida[col_idx - 1])
            celda.number_format = formato
            fila_salida[col_idx - 1] = celda
        ws.append(fila_salida)


@manejar_errores
def guardar_con_formato(path: Path, hojas: dict) -> bool:
    """
    hojas: {"NombreHoja": DataFrame, ...}. Si algun DataFrame viene vacio
    (0 filas) igual se escribe la hoja con encabezados, pero se avisa en
    consola y en el log.
    """
    hojas_vacias = [nombre for nombre, df in hojas.items() if len(df) == 0]

    wb = Workbook(write_only=True)
    for nombre_hoja, contenido in hojas.items():
        ws = wb.create_sheet(title=nombre_hoja)
        _formatear_hoja(ws, contenido)
    wb.save(path)

    print(f"[OK] Generado: {path}")
    for nombre_hoja in hojas_vacias:
        mensaje = f"   [!]  '{nombre_hoja}' en {path.name} quedo sin filas — revisa el filtro aplicado."
        print(mensaje)
        logger.warning(f"Hoja vacia: {nombre_hoja} en {path}")

    return True
